"""Reports as blocks over Odoo data, rendered once and exported to several formats.

Every number comes from a query run with the reader's own access rights; text
blocks are written by people (or pasted from a reviewed AI draft). Domains are
literal lists only (no code), fields must exist and be readable, and table rows
are capped.
"""

import ast
import io
import re

from markupsafe import Markup

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError
from odoo.tools import html2plaintext

MAX_ROWS = 500
BLOCK_TYPES = [('heading', 'Heading'), ('text', 'Text'), ('kpi', 'KPI tile'), ('table', 'Table')]
MEASURES = [('count', 'Number of records'), ('sum', 'Sum of a field'), ('avg', 'Average of a field')]


class Report(models.Model):
    _name = 'report.studio.report'
    _description = 'Report'
    _inherit = ['mail.thread']
    _order = 'write_date desc, id desc'

    name = fields.Char(required=True, tracking=True)
    subtitle = fields.Char()
    user_id = fields.Many2one('res.users', required=True, default=lambda self: self.env.user, string='Owner')
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company)
    shared = fields.Boolean(help='Everyone in the company can open, edit and export this report.')
    accent_color = fields.Char(default='#714B67', string='Accent colour')
    block_ids = fields.One2many('report.studio.block', 'report_id', copy=True, string='Blocks')
    preview_html = fields.Html(compute='_compute_preview', sanitize=False, string='Preview')
    res_model = fields.Char(readonly=True, help='Record this report was generated for, if any.')
    res_id = fields.Integer(readonly=True)

    @api.depends('name', 'subtitle', 'accent_color', 'block_ids', 'block_ids.sequence', 'block_ids.type',
                 'block_ids.title', 'block_ids.content', 'block_ids.model_id', 'block_ids.domain',
                 'block_ids.measure', 'block_ids.measure_field', 'block_ids.field_names', 'block_ids.limit')
    def _compute_preview(self):
        for report in self:
            try:
                styles = self.env['ir.qweb']._render('report_studio.report_styles', {})
                report.preview_html = styles + report._render_html()
            except (UserError, ValidationError) as error:
                report.preview_html = Markup('<div class="alert alert-warning">%s</div>') % str(error)

    # ------------------------------------------------------------------ data

    def _data(self):
        """Plain data for every block, shared by all renderers."""
        self.ensure_one()
        return {
            'title': self.name, 'subtitle': self.subtitle or '', 'accent': self.accent_color or '#714B67',
            'company': self.company_id.name, 'date': fields.Date.context_today(self),
            'blocks': [block._data() for block in self.block_ids.sorted('sequence')],
        }

    def _render_html(self):
        return self.env['ir.qweb']._render('report_studio.report_body', {'report': self._data()})

    # ------------------------------------------------------------------ exports

    def _export(self, fmt):
        self.ensure_one()
        self.check_access('read')
        data = self._data()
        stem = re.sub(r'[^\w-]+', '_', self.name).strip('_') or 'report'
        if fmt == 'html':
            page = self.env['ir.qweb']._render('report_studio.report_page', {'report': data, 'body': self._render_html()})
            return f'{stem}.html', 'text/html; charset=utf-8', str(page).encode()
        if fmt == 'pdf':
            content, _kind = self.env['ir.actions.report']._render_qweb_pdf('report_studio.action_report_pdf', self.ids)
            return f'{stem}.pdf', 'application/pdf', content
        if fmt == 'xlsx':
            return f'{stem}.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', _xlsx(data)
        if fmt == 'pptx':
            return f'{stem}.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation', _pptx(data)
        if fmt == 'docx':
            return f'{stem}.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', _docx(data)
        raise UserError(_('Unknown format: %s', fmt))

    def _download(self, fmt):
        self.ensure_one()
        return {'type': 'ir.actions.act_url', 'url': f'/report_studio/{self.id}/{fmt}', 'target': 'self'}

    def action_export_pdf(self):
        return self._download('pdf')

    def action_export_xlsx(self):
        return self._download('xlsx')

    def action_export_pptx(self):
        return self._download('pptx')

    def action_export_docx(self):
        return self._download('docx')

    def action_export_html(self):
        return self._download('html')


class Block(models.Model):
    _name = 'report.studio.block'
    _description = 'Report block'
    _order = 'sequence, id'

    report_id = fields.Many2one('report.studio.report', required=True, ondelete='cascade', index=True)
    sequence = fields.Integer(default=10)
    type = fields.Selection(BLOCK_TYPES, required=True, default='text')
    title = fields.Char()
    content = fields.Html(string='Text', sanitize=True)
    model_id = fields.Many2one('ir.model', string='Data from', ondelete='cascade',
                               domain=[('transient', '=', False)])
    model_name = fields.Char(related='model_id.model')
    domain = fields.Char(default='[]', help='Which records, e.g. [("state", "=", "sale")]. Literal values only.')
    measure = fields.Selection(MEASURES, default='count')
    measure_field = fields.Char(string='Field to add up', help='Technical name of a number field, e.g. amount_total.')
    field_names = fields.Char(string='Columns', help='Technical field names separated by commas, e.g. name, date, amount_total.')
    limit = fields.Integer(default=20, string='Rows')
    order = fields.Char(string='Sort by', help='e.g. date desc')

    @api.constrains('domain', 'model_id', 'field_names', 'measure_field', 'type', 'limit', 'order')
    def _check_query(self):
        for block in self.filtered(lambda b: b.type in ('kpi', 'table')):
            block._query_parts()

    def _query_parts(self):
        """Validated (model, domain, fields) for this block, read with the current user's rights."""
        if not self.model_id:
            raise ValidationError(_('Choose where the data of “%s” comes from.', self.title or self.type))
        source = self.sudo().model_id  # only the model's name; records are read with the user's rights
        if source.model not in self.env:
            raise ValidationError(_('The app behind “%s” is not installed.', self.title or self.type))
        Model = self.env[source.model]
        try:
            domain = ast.literal_eval(self.domain or '[]')
        except (ValueError, SyntaxError):
            raise ValidationError(_('The filter of “%s” must be a list such as [("state", "=", "done")].',
                                    self.title or self.type)) from None
        if not isinstance(domain, list):
            raise ValidationError(_('The filter must be a list.'))
        names = [name.strip() for name in (self.field_names or '').split(',') if name.strip()]
        for name in names + ([self.measure_field] if self.type == 'kpi' and self.measure != 'count' else []):
            if not name or name not in Model._fields:
                raise ValidationError(_('“%(field)s” is not a field of %(model)s.', field=name, model=source.name))
        if self.type == 'kpi' and self.measure != 'count' and Model._fields[self.measure_field].type not in (
                'integer', 'float', 'monetary'):
            raise ValidationError(_('Only number fields can be added up.'))
        if self.order and not re.fullmatch(r'\s*\w+(\s+(asc|desc))?\s*', self.order):
            raise ValidationError(_('Sort by one field, e.g. “date desc”.'))
        if not 0 < (self.limit or 0) <= MAX_ROWS:
            raise ValidationError(_('Tables show between 1 and %s rows.', MAX_ROWS))
        return Model, domain, names

    def _data(self):
        self.ensure_one()
        data = {'type': self.type, 'title': self.title or '', 'text': '', 'html': Markup(''),
                'value': None, 'columns': [], 'rows': []}
        if self.type == 'text':
            data['html'] = Markup(self.content or '')
            data['text'] = html2plaintext(self.content or '')
        elif self.type == 'kpi':
            Model, domain, _names = self._query_parts()
            if self.measure == 'count':
                value = Model.search_count(domain)
            else:
                aggregate = f'{self.measure_field}:{self.measure}'
                value = Model._read_group(domain, [], [aggregate])[0][0] or 0
            data['value'] = value
            data['display'] = f'{value:,.2f}' if isinstance(value, float) else f'{value:,}'
        elif self.type == 'table':
            Model, domain, names = self._query_parts()
            names = names or ['display_name']
            records = Model.search(domain, limit=self.limit, order=self.order or None)
            data['columns'] = [Model._fields[name].get_description(self.env)['string'] if name in Model._fields
                               else name for name in names]
            data['rows'] = [[_cell(record, name) for name in names] for record in records]
        return data


def _cell(record, name):
    value = record[name]
    field = record._fields.get(name)
    if field is None:
        return str(value)
    if field.type in ('many2one',):
        return value.display_name or ''
    if field.type in ('one2many', 'many2many'):
        return ', '.join(value.mapped('display_name'))
    if field.type == 'selection':
        return dict(field._description_selection(record.env)).get(value, value or '')
    if field.type == 'boolean':
        return '✓' if value else ''
    if value is False or value is None:
        return ''
    return value if isinstance(value, (int, float)) else str(value)


# ---------------------------------------------------------------------- office formats

def _xlsx(data):
    import xlsxwriter
    buffer = io.BytesIO()
    book = xlsxwriter.Workbook(buffer, {'in_memory': True})
    sheet = book.add_worksheet('Report')
    title = book.add_format({'bold': True, 'font_size': 16, 'font_color': data['accent']})
    heading = book.add_format({'bold': True, 'font_size': 12})
    head = book.add_format({'bold': True, 'bg_color': '#EEEEEE', 'border': 1})
    cell = book.add_format({'border': 1})
    row = 0
    sheet.write(row, 0, data['title'], title)
    row += 1
    if data['subtitle']:
        sheet.write(row, 0, data['subtitle'])
        row += 1
    row += 1
    for block in data['blocks']:
        if block['title']:
            sheet.write(row, 0, block['title'], heading)
            row += 1
        if block['type'] == 'text':
            for line in block['text'].splitlines() or ['']:
                sheet.write(row, 0, line)
                row += 1
        elif block['type'] == 'kpi':
            sheet.write(row, 0, block['value'])
            row += 1
        elif block['type'] == 'table':
            for col, name in enumerate(block['columns']):
                sheet.write(row, col, name, head)
                sheet.set_column(col, col, max(12, len(str(name)) + 2))
            row += 1
            for values in block['rows']:
                for col, value in enumerate(values):
                    sheet.write(row, col, value, cell)
                row += 1
        row += 1
    book.close()
    return buffer.getvalue()


def _missing(library):
    raise UserError(_('This export needs the Python package “%s” on the server. Ask your administrator to '
                      'install it (pip install %s).', library, library))


def _pptx(data):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        _missing('python-pptx')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = data['title']
    cover.placeholders[1].text = ' · '.join(filter(None, [data['subtitle'], data['company'], str(data['date'])]))
    kpis = [block for block in data['blocks'] if block['type'] == 'kpi']
    if kpis:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = 'Key figures'
        width = Inches(12.3 / max(len(kpis), 1))
        for index, block in enumerate(kpis[:6]):
            box = slide.shapes.add_textbox(Inches(0.5) + width * index, Inches(2.5), width, Inches(2)).text_frame
            box.text = block['display']
            box.paragraphs[0].runs[0].font.size = Pt(40)
            box.paragraphs[0].runs[0].font.bold = True
            para = box.add_paragraph()
            para.text = block['title']
            para.runs[0].font.size = Pt(16)
    for block in data['blocks']:
        if block['type'] == 'text':
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = block['title'] or data['title']
            slide.placeholders[1].text = block['text'][:2500]
        elif block['type'] == 'heading':
            slide = deck.slides.add_slide(deck.slide_layouts[2])
            slide.shapes.title.text = block['title']
        elif block['type'] == 'table' and block['columns']:
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = block['title'] or data['title']
            rows = block['rows'][:15]
            table = slide.shapes.add_table(len(rows) + 1, len(block['columns']), Inches(0.5), Inches(1.5),
                                           Inches(12.3), Inches(0.4) * (len(rows) + 1)).table
            for col, name in enumerate(block['columns']):
                table.cell(0, col).text = str(name)
            for r, values in enumerate(rows, 1):
                for col, value in enumerate(values):
                    table.cell(r, col).text = str(value)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _docx(data):
    try:
        import docx
    except ImportError:
        _missing('python-docx')
    document = docx.Document()
    document.add_heading(data['title'], 0)
    if data['subtitle']:
        document.add_paragraph(data['subtitle'])
    document.add_paragraph(f"{data['company']} · {data['date']}")
    for block in data['blocks']:
        if block['type'] == 'heading':
            document.add_heading(block['title'], 1)
            continue
        if block['title']:
            document.add_heading(block['title'], 2)
        if block['type'] == 'text':
            for line in block['text'].split('\n\n'):
                document.add_paragraph(line.strip())
        elif block['type'] == 'kpi':
            document.add_paragraph().add_run(block['display']).bold = True
        elif block['type'] == 'table' and block['columns']:
            table = document.add_table(rows=1, cols=len(block['columns']))
            table.style = 'Light Grid Accent 1'
            for col, name in enumerate(block['columns']):
                table.rows[0].cells[col].text = str(name)
            for values in block['rows']:
                cells = table.add_row().cells
                for col, value in enumerate(values):
                    cells[col].text = str(value)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
